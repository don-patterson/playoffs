"""
Extract playoff picks from Cup Bracket 2026.xlsx.

Each person has their own sheet. Picks are recorded as:
  - text like "Avs in 6" or "Wild 7"  (team + game count)
  - a bare number like 6              (game count only; team comes from adjacent column)

Outputs:
  - `picks`     : dict  {person -> {matchup_id -> {description, teams, pick, games}}}
  - picks.csv   : CSV with one row per person per matchup
"""

import csv
import json
import re

import openpyxl

# ---------------------------------------------------------------------------
# Team name normalisation
# ---------------------------------------------------------------------------

TEAM_ALIASES: dict[str, str] = {
    # Avalanche
    "avs": "Avalanche", "abs": "Avalanche", "avalanche": "Avalanche",
    # Kings
    "kings": "Kings",
    # Stars
    "stars": "Stars",
    # Wild
    "wild": "Wild",
    # Golden Knights
    "gk": "Golden Knights", "knights": "Golden Knights",
    "kinghts": "Golden Knights",          # common typo
    "golden knights": "Golden Knights",
    # Mammoth
    "mammoth": "Mammoth",
    # Oilers
    "oil": "Oilers", "oilers": "Oilers",
    # Ducks
    "ducks": "Ducks", "duck": "Ducks", "anaheim": "Ducks",
    # Sabres / Buffalo
    "sabres": "Buffalo", "sabers": "Buffalo", "sabs": "Buffalo",
    "buffalo": "Buffalo",
    # Bruins
    "bruins": "Bruins", "boston": "Bruins",
    # Lightning
    "lightning": "Lightning", "bolts": "Lightning",
    # Canadiens
    "canadiens": "Canadiens", "habs": "Canadiens", "montreal": "Canadiens",
    # Hurricanes
    "hurricanes": "Hurricanes", "canes": "Hurricanes",
    "hurricane": "Hurricanes", "carolina": "Hurricanes",
    # Senators
    "senators": "Senators", "sens": "Senators", "ottawa": "Senators",
    # Penguins
    "penguins": "Penguins", "pens": "Penguins", "pittsburgh": "Penguins",
    # Flyers
    "flyers": "Flyers", "philadelphia": "Flyers",
}

HEADER_VALUES = {
    "RD 1 Picks", "RD 2 Picks", "RD 3 Picks",
    "Cup Pick", "Point Cats", "Points",
    "West", "East", "Stanley Cup Final",
}


def normalize_team(name: str | None) -> str | None:
    if not name:
        return None
    cleaned = str(name).strip()
    return TEAM_ALIASES.get(cleaned.lower(), cleaned)


# ---------------------------------------------------------------------------
# Pick parsing helpers
# ---------------------------------------------------------------------------

def parse_text_pick(text: str) -> tuple[str | None, int | None]:
    """Return (team_str, games) from a free-text pick such as 'Avs in 6'."""
    text = text.strip()
    # "Team in N"  (with or without space before the digit)
    m = re.match(r"(.+?)\s+in\s*(\d+)\s*$", text, re.IGNORECASE)
    if m:
        return m.group(1).strip(), int(m.group(2))
    # "Team N"
    m = re.match(r"(.+?)\s+(\d+)\s*$", text)
    if m:
        return m.group(1).strip(), int(m.group(2))
    # "Team in ..."  (no game count entered)
    m = re.match(r"(.+?)\s+in\s*\.+\s*$", text, re.IGNORECASE)
    if m:
        return m.group(1).strip(), None
    return text, None


def get_pick(
    ws,
    pick_cells: list[str],
    team_col_map: dict[str, str],
) -> tuple[str | None, int | None]:
    """
    Scan *pick_cells* in order and return (team, games) for the first valid
    pick found.  *team_col_map* maps pick-column letter → team-column letter
    (used when the pick is stored as a bare number).
    """
    for addr in pick_cells:
        cell = ws[addr]
        value = cell.value
        if value is None:
            continue
        if isinstance(value, str) and value.strip() in HEADER_VALUES:
            continue

        row = cell.row
        col = cell.column_letter

        if isinstance(value, (int, float)):
            team_col = team_col_map.get(col)
            team_hint = ws[f"{team_col}{row}"].value if team_col else None
            return normalize_team(team_hint), int(value)

        # String pick
        team_str, games = parse_text_pick(str(value))
        return normalize_team(team_str), games

    return None, None


# ---------------------------------------------------------------------------
# Bracket structure
# Each entry: (id, description, team1_cell, team2_cell, pick_cells, team_col_map)
#
# team1_cell / team2_cell  – canonical team cells visible on every sheet
# pick_cells               – cells searched in order for the actual pick
# team_col_map             – {pick_col_letter: team_col_letter} for bare-number picks
# ---------------------------------------------------------------------------

MATCHUPS: list[tuple] = [
    # ── West Round 1 ──────────────────────────────────────────────────────
    ("west_r1_avs_kings",
     "West R1: Avalanche vs Kings",
     "A3", "A4",
     ["B3", "B4"], {"B": "A"}),

    ("west_r1_stars_wild",
     "West R1: Stars vs Wild",
     "A6", "A7",
     ["B6", "B7"], {"B": "A"}),

    ("west_r1_gk_mammoth",
     "West R1: Golden Knights vs Mammoth",
     "A9", "A10",
     ["B9", "B10"], {"B": "A"}),

    ("west_r1_oilers_ducks",
     "West R1: Oilers vs Ducks",
     "A12", "A13",
     ["B12", "B13"], {"B": "A"}),

    # ── West Round 2 ──────────────────────────────────────────────────────
    ("west_r2_upper",
     "West R2 (Avs/Kings vs Stars/Wild winner)",
     "C5", "C6",
     ["D5", "D6"], {"D": "C"}),

    ("west_r2_lower",
     "West R2 (GK/Mammoth vs Oilers/Ducks winner)",
     "C10", "C11",
     ["D9", "D10", "D11"], {"D": "C"}),

    # ── West Conference Final ──────────────────────────────────────────────
    ("west_r3",
     "West Conference Final",
     "E7", "E8",
     ["F7", "F8"], {"F": "E"}),

    # ── East Round 1 ──────────────────────────────────────────────────────
    # Note: B16 holds the "RD 1 Picks" column header on the East section,
    # so picks start at B17.  B18 is included because some players put the
    # Lightning/Canadiens pick there (in sheets where B17 already has the
    # Buffalo/Bruins pick).
    ("east_r1_buf_bos",
     "East R1: Buffalo vs Bruins",
     "A16", "A17",
     ["B17", "B18"], {"B": "A"}),

    ("east_r1_tbl_mtl",
     "East R1: Lightning vs Canadiens",
     "A19", "A20",
     ["B18", "B19", "B20"], {"B": "A"}),

    ("east_r1_car_ott",
     "East R1: Hurricanes vs Senators",
     "A22", "A23",
     ["B22", "B23", "B24"], {"B": "A"}),

    ("east_r1_pit_phi",
     "East R1: Penguins vs Flyers",
     "A25", "A26",
     ["B25", "B26"], {"B": "A"}),

    # ── East Round 2 ──────────────────────────────────────────────────────
    ("east_r2_upper",
     "East R2 (Buf/Bos vs TBL/MTL winner)",
     "C19", "C20",
     ["D19", "D20"], {"D": "C"}),

    ("east_r2_lower",
     "East R2 (CAR/OTT vs PIT/PHI winner)",
     "C24", "C25",
     ["D23", "D24", "D25"], {"D": "C"}),

    # ── East Conference Final ──────────────────────────────────────────────
    ("east_r3",
     "East Conference Final",
     "E21", "E22",
     ["F20", "F21"], {"F": "E"}),

    # ── Stanley Cup Final ──────────────────────────────────────────────────
    ("cup",
     "Stanley Cup Final",
     "G14", "G15",
     ["H12", "H13", "H14", "H15"], {"H": "G"}),
]


# ---------------------------------------------------------------------------
# Main extraction
# ---------------------------------------------------------------------------

def extract_picks(
    filepath: str = "Cup Bracket 2026.xlsx",
    skip_sheets: tuple[str, ...] = ("Demo",),
) -> dict:
    wb = openpyxl.load_workbook(filepath)
    all_picks: dict = {}

    for sheet_name in wb.sheetnames:
        if sheet_name in skip_sheets:
            continue

        ws = wb[sheet_name]
        person_picks: dict = {}

        for matchup_id, description, t1_cell, t2_cell, pick_cells, team_col_map in MATCHUPS:
            team1 = normalize_team(ws[t1_cell].value)
            team2 = normalize_team(ws[t2_cell].value)

            # For R1 the teams are fixed in the template; for later rounds they
            # are filled in by the participant as they advance teams.
            teams = (team1, team2) if (team1 or team2) else None

            pick_team, games = get_pick(ws, pick_cells, team_col_map)

            person_picks[matchup_id] = {
                "description": description,
                "teams": teams,
                "pick": pick_team,
                "games": games,
            }

        all_picks[sheet_name] = person_picks

    return all_picks


# ---------------------------------------------------------------------------
# Output helpers
# ---------------------------------------------------------------------------

def write_csv(picks: dict, out_path: str = "picks.csv") -> None:
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["person", "matchup_id", "description",
                         "team1", "team2", "pick", "games"])
        for person, matchups in picks.items():
            for mid, data in matchups.items():
                teams = data.get("teams") or ("", "")
                writer.writerow([
                    person,
                    mid,
                    data["description"],
                    teams[0] or "",
                    teams[1] or "",
                    data.get("pick") or "",
                    data.get("games") if data.get("games") is not None else "",
                ])


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    picks = extract_picks()

    # Pretty-print the dictionary
    print(json.dumps(picks, indent=2, ensure_ascii=False))

    # Write CSV
    write_csv(picks)
    print("\n→ picks.csv written")
